from pathlib import Path
import textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image, ImageOps

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'presentations'
W,H=13.333,7.5
BG='#0D1424'; CARD='#182237'; WHITE='#F5F7FB'; MUTED='#B9C3D4'; AMBER='#FFB021'; LINE='#2C3A52'
FONT='/Library/Fonts/Arial Unicode.ttf'
pdfmetrics.registerFont(TTFont('Deck',FONT))
prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
pdf=canvas.Canvas(str(OUT/'ServeOS-Product-Brief.pdf'),pagesize=(W*72,H*72))
pdf.setTitle('ServeOS — Product Brief'); pdf.setAuthor('ServeOS')
slide=None

def rgb(color): return RGBColor.from_string(color.lstrip('#'))
def rect(x,y,w,h,color=CARD,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=rgb(color);shape.line.fill.background()
    if radius:
        try: shape.adjustments[0]=0.12
        except Exception: pass
    pdf.setFillColor(color)
    if radius:pdf.roundRect(x*72,(H-y-h)*72,w*72,h*72,10,stroke=0,fill=1)
    else:pdf.rect(x*72,(H-y-h)*72,w*72,h*72,stroke=0,fill=1)

def txt(x,y,w,h,text,size=18,color=WHITE,bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame;tf.clear();tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    tf.word_wrap=False
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(0);p.line_spacing=1.12
    pdf.setFillColor(color);pdf.setFont('Deck',size)
    baseline=(H-y)*72-size*0.89
    for line in text.split('\n'):
        pdf.drawString(x*72,baseline,line);baseline-=size*1.12

def image(x,y,w,h,path,name):
    im=Image.open(ROOT/path).convert('RGB');im=ImageOps.fit(im,(int(w*130),int(h*130)),method=Image.Resampling.LANCZOS)
    cropped=OUT/f'{name}.jpg';im.save(cropped,quality=92)
    pic=slide.shapes.add_picture(str(cropped), Inches(x), Inches(y),width=Inches(w),height=Inches(h));pic.name=name
    pdf.drawImage(ImageReader(im),x*72,(H-y-h)*72,w*72,h*72)

def new(n,label,title=None,subtitle=None):
    global slide
    if n>1:pdf.showPage()
    slide=prs.slides.add_slide(prs.slide_layouts[6]);rect(0,0,W,H,BG)
    txt(.65,.36,6,.28,'ServeOS',16,WHITE,True);txt(.65,6.98,10,.2,'PRODUCT BRIEF  /  14 OCT 2026',8,MUTED)
    txt(12.0,6.98,.7,.2,f'{n:02d} / 07',8,MUTED)
    rect(.65,.86,.45,.035,AMBER)
    txt(.65,1.10,11,.25,label.upper(),10,AMBER,True)
    if title:txt(.65,1.55,12,.76,title,32,WHITE,True)
    if subtitle:txt(.65,2.32,12,.50,subtitle,16,MUTED)

def card(x,y,w,title,body,num=None):
    rect(x,y,w,2.5,CARD,True)
    if num:txt(x+.25,y+.24,w-.5,.35,num,13,AMBER,True)
    txt(x+.25,y+.78,w-.5,.6,title,20,WHITE,True)
    txt(x+.25,y+1.45,w-.5,.9,body,15,MUTED)

new(1,'Restaurant management, connected')
image(7.52,0,5.82,6.75,'src/assets/hero-food.jpg','cover-food')
rect(7.52,6.05,5.82,.7,CARD)
txt(7.86,6.25,5,.3,'WEBSITE  •  ORDERS  •  OPERATIONS',11,AMBER,True)
txt(.65,1.70,6.6,2.2,'One platform.\nEvery restaurant\nworkflow.',40,WHITE,True)
txt(.65,4.12,6.25,1.0,'Restaurant websites, QR ordering,\nteam workflows and AI support.',21,MUTED)
rect(.65,5.61,4.9,.64,AMBER,True);txt(.93,5.80,4.45,.3,'Start with a 7-day free trial',17,BG,True)
slide.notes_slide.notes_text_frame.text='ServeOS is a SaaS platform for restaurant owners and the teams serving their guests. This brief covers the current product and packages.'

new(2,'The opportunity','Bring scattered work into one flow.','A connected experience for the guest, the restaurant team and the platform operator.')
card(.65,3.14,3.88,'Be visible online','Publish a restaurant website\nand keep the digital menu\neasy to browse.','01 / DISCOVERY')
card(4.73,3.14,3.88,'Keep service moving','Connect table ordering,\nwaiter service and\nkitchen tickets.','02 / SERVICE')
card(8.81,3.14,3.87,'Stay in control','Manage staff, subscriptions,\nbranding and business tools\nfrom dedicated panels.','03 / MANAGEMENT')
slide.notes_slide.notes_text_frame.text='The product is designed to connect restaurant discovery, ordering and daily operations. We make no quantified efficiency or revenue claims.'

new(3,'Guest experience','A restaurant website that can take orders.')
image(.65,2.73,7.25,3.73,'public/demo-screenshots/guest-homepage.png','guest-preview')
for y,title,body in [(2.88,'Your public presence','Restaurant pages, branding and themes.'),(3.93,'Your digital menu','Categories, prices, images and availability.'),(4.98,'Your table ordering','QR links connect guests to the right menu.')]:
    txt(8.24,y,4.5,.35,title,20,WHITE,True);txt(8.24,y+.45,4.45,.6,body,14,MUTED)
txt(.65,6.60,7,.2,'PRODUCT INTERFACE PREVIEW',8,MUTED)
slide.notes_slide.notes_text_frame.text='Guests can discover a restaurant and browse its image-rich digital menu. Table QR ordering connects the guest to the appropriate restaurant and table. Preview imagery comes from project demo assets.'

new(4,'Restaurant operations','Give every team member the right workspace.')
image(.65,2.75,7.25,3.76,'public/demo-screenshots/kitchen-display.png','kitchen-preview')
for y,title,body in [(2.82,'Owner & admin','Menus, staff, orders, billing and settings.'),(3.84,'Kitchen & waiter','Dedicated panels for tickets and service.'),(4.86,'Business tools','Accounting, loyalty, promotions and reports.')]:
    txt(8.24,y,4.5,.36,title,20,WHITE,True);txt(8.24,y+.47,4.4,.55,body,14,MUTED)
txt(.65,6.60,7,.2,'PRODUCT INTERFACE PREVIEW  /  FEATURES DEPEND ON PACKAGE AND ROLE',8,MUTED)
slide.notes_slide.notes_text_frame.text='Restaurant owners and admins create staff accounts with specific roles. Modules and capacity depend on the subscription, role and active add-ons.'

new(5,'Support & SaaS control','ServeOS AI, backed by a human team.')
rect(.65,2.77,6.05,3.81,CARD,True);txt(.95,3.03,5.55,.4,'Customer support',21,WHITE,True)
for i,(title,body) in enumerate([('01  Customer message','A question from website or restaurant chat.'),('02  ServeOS AI','Customer-language replies using published guides.'),('03  Answer or human handoff','Admins can review drafts and take over the chat.')]):
    y=3.79+i*.80;txt(.95,y,5.5,.28,title,15,AMBER,True);txt(.95,y+.34,5.45,.30,body,13,MUTED)
rect(6.93,2.77,5.76,3.81,CARD,True);txt(7.23,3.03,5.2,.4,'Platform administration',21,WHITE,True)
txt(7.23,3.84,5.2,2.2,'Restaurant tenants and packages\nPayment review and add-on catalog\nWebsite themes and domains\nSupport inbox and AI knowledge\nGemini API connection settings',17,MUTED)
txt(.65,6.68,12,.2,'AI REPLIES REQUIRE AN AVAILABLE, CONFIGURED GEMINI CONNECTION.',8,MUTED)
slide.notes_slide.notes_text_frame.text='ServeOS AI supports automatic replies and reviewed drafts using an editable platform knowledge base. Human requests and unsupported issues can be handed to the support team. Gemini service availability is required. Platform admins control tenants, packages, payment review, themes, domains and AI settings.'

new(6,'Packages','Start free. Choose the term that fits.','Current platform package prices in Bangladeshi taka (BDT).')
rect(.65,3.05,12.04,.61,AMBER,True);txt(.96,3.22,11.6,.29,'7-DAY FREE TRIAL   /   New restaurants   /   No automatic charge',17,BG,True)
packages=[('MONTHLY','349','1 month','Try the platform'),('3 MONTHS','944','3 months','Most popular'),('6 MONTHS','1,444','6 months','Longer commitment'),('YEARLY','2,197','12 months','Best value')]
for i,(name,price,term,tag) in enumerate(packages):
    x=.65+i*3.07;rect(x,3.97,2.84,2.54,CARD,True);txt(x+.22,4.24,2.4,.3,name,12,AMBER,True);txt(x+.22,4.91,2.4,.6,price,33,WHITE,True);txt(x+.22,5.60,2.4,.3,f'BDT total / {term}',12,MUTED);txt(x+.22,6.06,2.4,.28,tag,11,MUTED)
slide.notes_slide.notes_text_frame.text='Prices were checked against the local public package catalog on 14 October 2026. Amounts are totals for the stated term, not monthly equivalents. The seven-day trial applies to new restaurants. Prices and features are editable by platform administrators.'

new(7,'Getting started','From signup to your first table scan.')
for x,num,title,body in [(.65,'01','Create your restaurant','Sign up and complete\nthe restaurant details.'),(4.73,'02','Set up your team & menu','Add dishes, photos,\ntables and staff accounts.'),(8.81,'03','Share your restaurant','Share the public website\nand test your table QR links.')]:
    card(x,2.89,3.88,title,body,num)
rect(.65,5.85,12.04,.72,AMBER,True);txt(.98,6.07,11.5,.34,'ServeOS  /  One platform for the restaurant and the team behind it.',18,BG,True)
slide.notes_slide.notes_text_frame.text='Start with the free trial, set up the restaurant menu and staff, then test the guest ordering experience. ServeOS is the product name; the demo remains a local application rather than a publicly hosted deployment.'

prs.core_properties.title='ServeOS — Product Brief';prs.core_properties.subject='Restaurant SaaS product overview';prs.core_properties.author='ServeOS'
prs.save(OUT/'ServeOS-Product-Brief.pptx');pdf.save()
for path in OUT.glob('*-preview.jpg'): path.unlink()
# Cropped images are embedded in both deliverables; keep the folder tidy.
for name in ['cover-food.jpg','guest-preview.jpg','kitchen-preview.jpg']:
    (OUT/name).unlink(missing_ok=True)
print('Created 7-slide PowerPoint and PDF in presentations/.')
